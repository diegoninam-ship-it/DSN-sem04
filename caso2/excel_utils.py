import os

from openpyxl import Workbook, load_workbook

EXCEL_PATH = os.path.join(os.path.dirname(__file__), "miembros_mesa.xlsx")

ENCABEZADOS = ["DNI", "Región", "Provincia", "Distrito", "Dirección del local de votación"]


def inicializar_excel():
    """Crea el archivo con encabezados si todavía no existe."""
    if not os.path.exists(EXCEL_PATH):
        wb = Workbook()
        hoja = wb.active
        hoja.title = "Miembros de Mesa"
        hoja.append(ENCABEZADOS)
        wb.save(EXCEL_PATH)


def agregar_registro(dni, region, provincia, distrito, direccion):
    """Agrega una fila nueva al Excel."""
    inicializar_excel()
    wb = load_workbook(EXCEL_PATH)
    hoja = wb.active
    hoja.append([dni, region, provincia, distrito, direccion])
    wb.save(EXCEL_PATH)


def listar_registros():
    """Devuelve todas las filas ya guardadas (sin encabezados), más recientes primero."""
    inicializar_excel()
    wb = load_workbook(EXCEL_PATH)
    hoja = wb.active
    filas = list(hoja.iter_rows(min_row=2, values_only=True))
    filas.reverse()
    return filas